#!/usr/bin/env python3
"""
Generate the Official RailGaadi PowerPoint Presentation (.pptx)
Matching the exact layout, dimensions, fonts, aspect ratio, college logos,
and white-background/black-text color scheme of the original Poornima ITS PPT.

Author: Nikhil Kumar (PCE23IT034)
Session: 2026-27
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(output_path):
    prs = Presentation()
    # 16:9 widescreen dimensions matching the original file
    prs.slide_width = Inches(10.0)
    prs.slide_height = Inches(5.625)

    blank_layout = prs.slide_layouts[6]

    # Colors: White background, Black text, Gray divider line
    C_WHITE = RGBColor(255, 255, 255)
    C_BLACK = RGBColor(0, 0, 0)
    C_GRAY_LINE = RGBColor(158, 158, 158)  # #9E9E9E from original PPT

    logo_slide1 = os.path.abspath("public/presentation_assets/slide1_logo.png")
    logo_header = os.path.abspath("public/presentation_assets/slide_header_logo.png")

    snap_map = os.path.abspath("public/presentation_assets/snap_map.png")
    snap_board = os.path.abspath("public/presentation_assets/snap_board.png")
    snap_analytics = os.path.abspath("public/presentation_assets/snap_analytics.png")
    snap_terminal = os.path.abspath("public/presentation_assets/snap_terminal.png")

    def set_white_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10.0), Inches(5.625))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_WHITE
        bg.line.color.rgb = C_WHITE
        return bg

    def add_standard_header(slide, title_text):
        # 1. Slide Title
        t_box = slide.shapes.add_textbox(Inches(0.57), Inches(0.12), Inches(8.30), Inches(0.87))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Calibri"
        p.font.size = Pt(38)
        p.font.bold = False
        p.font.color.rgb = C_BLACK

        # 3. Horizontal Divider Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(1.11), Inches(10.0), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = C_GRAY_LINE
        line.line.color.rgb = C_GRAY_LINE

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_white_bg(s1)

    # Header Box
    tb_head = s1.shapes.add_textbox(Inches(0.54), Inches(0.20), Inches(8.91), Inches(1.20))
    tf_head = tb_head.text_frame
    tf_head.word_wrap = True
    tf_head.margin_left = tf_head.margin_top = tf_head.margin_right = tf_head.margin_bottom = 0
    
    p1 = tf_head.paragraphs[0]
    p1.text = "POORNIMA COLLEGE OF ENGINEERING"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = "Calibri"
    p1.font.size = Pt(36)
    p1.font.bold = False
    p1.font.color.rgb = C_BLACK

    p2 = tf_head.add_paragraph()
    p2.text = "DEPARTMENT OF INFORMATION TECHNOLOGY"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Calibri"
    p2.font.size = Pt(22)
    p2.font.bold = False
    p2.font.color.rgb = C_BLACK
    p2.space_before = Pt(4)

    # Presentation On Box
    tb_sub = s1.shapes.add_textbox(Inches(1.5), Inches(1.35), Inches(7.0), Inches(1.50))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0

    p_an = tf_sub.paragraphs[0]
    p_an.text = "An"
    p_an.alignment = PP_ALIGN.CENTER
    p_an.font.name = "Calibri"
    p_an.font.size = Pt(18)
    p_an.font.color.rgb = C_BLACK

    p_its = tf_sub.add_paragraph()
    p_its.text = "INDUSTRIAL TRAINING SEMINAR PRESENTATION"
    p_its.alignment = PP_ALIGN.CENTER
    p_its.font.name = "Calibri"
    p_its.font.size = Pt(18)
    p_its.font.color.rgb = C_BLACK
    p_its.space_before = Pt(2)

    p_on = tf_sub.add_paragraph()
    p_on.text = "on"
    p_on.alignment = PP_ALIGN.CENTER
    p_on.font.name = "Calibri"
    p_on.font.size = Pt(18)
    p_on.font.color.rgb = C_BLACK
    p_on.space_before = Pt(2)

    p_title = tf_sub.add_paragraph()
    p_title.text = "Software Development"
    p_title.alignment = PP_ALIGN.CENTER
    p_title.font.name = "Calibri"
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = C_BLACK
    p_title.space_before = Pt(4)

    # Poornima College Logo (Centered)
    if os.path.exists(logo_slide1):
        s1.shapes.add_picture(logo_slide1, Inches(4.28), Inches(2.92), Inches(1.44), Inches(1.44))

    # Bottom Details: Presented to (Left)
    tb_left = s1.shapes.add_textbox(Inches(0.54), Inches(4.25), Inches(3.40), Inches(1.30))
    tf_left = tb_left.text_frame
    tf_left.word_wrap = True
    tf_left.margin_left = tf_left.margin_top = tf_left.margin_right = tf_left.margin_bottom = 0

    p_to = tf_left.paragraphs[0]
    p_to.text = "Presented to:"
    p_to.font.name = "Calibri"
    p_to.font.size = Pt(16.5)
    p_to.font.color.rgb = C_BLACK

    p_c1 = tf_left.add_paragraph()
    p_c1.text = "Mr. Gajendra Singh Rajawat"
    p_c1.font.name = "Calibri"
    p_c1.font.size = Pt(16.5)
    p_c1.font.bold = True
    p_c1.font.color.rgb = C_BLACK
    p_c1.space_before = Pt(2)

    p_c2 = tf_left.add_paragraph()
    p_c2.text = "Ms. Priya Shekhawat"
    p_c2.font.name = "Calibri"
    p_c2.font.size = Pt(16.5)
    p_c2.font.bold = True
    p_c2.font.color.rgb = C_BLACK

    p_c3 = tf_left.add_paragraph()
    p_c3.text = "(Faculty Coordinators)"
    p_c3.font.name = "Calibri"
    p_c3.font.size = Pt(16.5)
    p_c3.font.color.rgb = C_BLACK

    # Bottom Details: Session (Center)
    tb_mid = s1.shapes.add_textbox(Inches(4.00), Inches(4.45), Inches(2.00), Inches(0.50))
    tf_mid = tb_mid.text_frame
    tf_mid.margin_left = tf_mid.margin_top = tf_mid.margin_right = tf_mid.margin_bottom = 0
    p_sess = tf_mid.paragraphs[0]
    p_sess.text = "Session: 2026-27"
    p_sess.alignment = PP_ALIGN.CENTER
    p_sess.font.name = "Calibri"
    p_sess.font.size = Pt(16.5)
    p_sess.font.color.rgb = C_BLACK

    # Bottom Details: Presented By (Right)
    tb_right = s1.shapes.add_textbox(Inches(7.60), Inches(4.30), Inches(2.00), Inches(1.10))
    tf_right = tb_right.text_frame
    tf_right.word_wrap = True
    tf_right.margin_left = tf_right.margin_top = tf_right.margin_right = tf_right.margin_bottom = 0

    p_by = tf_right.paragraphs[0]
    p_by.text = "Presented By:"
    p_by.font.name = "Calibri"
    p_by.font.size = Pt(16.5)
    p_by.font.color.rgb = C_BLACK

    p_name = tf_right.add_paragraph()
    p_name.text = "Nikhil Kumar"
    p_name.font.name = "Calibri"
    p_name.font.size = Pt(16.5)
    p_name.font.bold = True
    p_name.font.color.rgb = C_BLACK
    p_name.space_before = Pt(2)

    p_reg = tf_right.add_paragraph()
    p_reg.text = "PCE23IT034"
    p_reg.font.name = "Calibri"
    p_reg.font.size = Pt(16.5)
    p_reg.font.bold = True
    p_reg.font.color.rgb = C_BLACK

    # =========================================================================
    # SLIDE 2: CONTENT
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_white_bg(s2)
    add_standard_header(s2, "CONTENT")

    tb2 = s2.shapes.add_textbox(Inches(0.57), Inches(1.25), Inches(8.90), Inches(4.20))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0

    content_items = [
        ("•   Training Platform", False),
        ("•   Why you have selected this platform?", False),
        ("•   Prerequisites of the Training", False),
        ("•   Week wise schedule", False),
        ("•   Technology Used", False),
        ("•   Project Description", False),
        ("       –   Project Introduction", False),
        ("       –   Project Scope", False),
        ("       –   Module Assigned", False),
        ("•   Learning", False),
        ("•   Snapshots of Project", False),
        ("•   Important Code Snippets", False),
        ("•   Future Scope", False),
        ("•   Conclusion", False),
        ("•   References", False)
    ]

    for idx, (text, is_bold) in enumerate(content_items):
        p = tf2.paragraphs[0] if idx == 0 else tf2.add_paragraph()
        p.text = text
        p.font.name = "Calibri"
        p.font.size = Pt(14.5)
        p.font.bold = is_bold
        p.font.color.rgb = C_BLACK
        p.space_before = Pt(2.5)

    # =========================================================================
    # SLIDE 3: Training Platform
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_white_bg(s3)
    add_standard_header(s3, "Training Platform")

    tb3 = s3.shapes.add_textbox(Inches(0.57), Inches(1.30), Inches(8.90), Inches(4.10))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = tf3.margin_top = tf3.margin_right = tf3.margin_bottom = 0

    p_sub3 = tf3.paragraphs[0]
    p_sub3.text = "About Sumit DigiTech Private Limited"
    p_sub3.font.name = "Calibri"
    p_sub3.font.size = Pt(18)
    p_sub3.font.bold = True
    p_sub3.font.color.rgb = C_BLACK
    p_sub3.space_after = Pt(14)

    s3_bullets = [
        [("●   Industry: ", True), ("Information Technology & Digital Solutions | ", False), ("Location: ", True), ("India", False)],
        [("●   Services: ", True), ("Software Development, Web Development, Website Design, UI/UX Development, Digital Solutions", False)],
        [("●   Focus Areas: ", True), ("Modern Web Applications, Software Solutions, Digital Transformation, Technology-Driven Development", False)],
        [("●   Training & Projects: ", True), ("Practical software development projects and hands-on technical learning opportunities", False)],
        [("●   Development Approach: ", True), ("User-Focused Solutions • Quality Development • Modern Technologies • Practical Implementation", False)],
        [("●   Core Values: ", True), ("Innovation • Quality • Customer Satisfaction • Reliable & User-Friendly Solutions", False)]
    ]

    for item in s3_bullets:
        p = tf3.add_paragraph()
        p.space_before = Pt(10)
        for text, bold in item:
            r = p.add_run()
            r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(14.5)
            r.font.bold = bold
            r.font.color.rgb = C_BLACK

    # =========================================================================
    # SLIDE 4: Why I Selected This Platform
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_white_bg(s4)
    add_standard_header(s4, "Why I Selected This Platform")

    tb4 = s4.shapes.add_textbox(Inches(0.57), Inches(1.35), Inches(8.90), Inches(4.00))
    tf4 = tb4.text_frame
    tf4.word_wrap = True
    tf4.margin_left = tf4.margin_top = tf4.margin_right = tf4.margin_bottom = 0

    s4_bullets = [
        [("●   Practical Learning Opportunities – ", True), ("Exposure to Software Development, Web Development, and real-time transit telemetry applications.", False)],
        [("●   Professional Work Culture – ", True), ("Supportive and learning-oriented environment focused on innovation and clean software architecture.", False)],
        [("●   Real-World Experience – ", True), ("Hands-on involvement in developing practical software projects, GIS cartography, and scalable web solutions.", False)],
        [("●   Skill Development – ", True), ("Improved technical knowledge, spatial mathematical algorithms, frontend performance, and problem-solving abilities.", False)],
        [("●   Growth-Oriented – ", True), ("Guidance from experienced professionals, continuous code reviews, and valuable teamwork experience.", False)],
        [("●   Project-Based Learning – ", True), ("Opportunity to design and build 'RailGaadi' — a live passenger train tracking and route intelligence web platform.", False)],
        [("●   Future Relevance – ", True), ("Technologies and development practices aligned with current and emerging industry trends in modern web standards.", False)]
    ]

    for idx, item in enumerate(s4_bullets):
        p = tf4.paragraphs[0] if idx == 0 else tf4.add_paragraph()
        p.space_before = Pt(10)
        for text, bold in item:
            r = p.add_run()
            r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(14.5)
            r.font.bold = bold
            r.font.color.rgb = C_BLACK

    # =========================================================================
    # SLIDE 5: Prerequisites of the Training
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_white_bg(s5)
    add_standard_header(s5, "Prerequisites of the Training")

    tb5 = s5.shapes.add_textbox(Inches(0.57), Inches(1.35), Inches(8.90), Inches(4.00))
    tf5 = tb5.text_frame
    tf5.word_wrap = True
    tf5.margin_left = tf5.margin_top = tf5.margin_right = tf5.margin_bottom = 0

    s5_bullets = [
        [("●   Fundamentals of Web Technologies – ", True), ("Strong understanding of semantic HTML5, modern CSS3, and JavaScript logic.", False)],
        [("●   Frontend Development – ", True), ("Practical knowledge of modern ES Modules, component-based layout design, and responsive user interfaces.", False)],
        [("●   State Architecture & Asynchronous Logic – ", True), ("Understanding event-driven Pub-Sub reactive patterns, Promises, and REST API communication.", False)],
        [("●   Spatial Cartography & GIS Math – ", True), ("Basic knowledge of geographical coordinate systems (Lat/Lon), Haversine spherical math, and vector maps.", False)],
        [("●   API Integration & Reverse Proxies – ", True), ("Practical exposure to transit APIs, OpenWeatherMap, Vite dev proxy, and CORS handling.", False)],
        [("●   Version Control & Development Tools – ", True), ("Using Git/GitHub, VS Code, npm, Vite bundler, and browser DevTools for debugging.", False)],
        [("●   Problem-Solving & Soft Skills – ", True), ("Debugging ability, logical thinking, communication, teamwork, and technical adaptability.", False)]
    ]

    for idx, item in enumerate(s5_bullets):
        p = tf5.paragraphs[0] if idx == 0 else tf5.add_paragraph()
        p.space_before = Pt(10)
        for text, bold in item:
            r = p.add_run()
            r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(14.5)
            r.font.bold = bold
            r.font.color.rgb = C_BLACK

    # =========================================================================
    # SLIDE 6: Week Wise Schedule (Weeks 1 to 3)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_white_bg(s6)
    add_standard_header(s6, "Week Wise Schedule")

    tb6 = s6.shapes.add_textbox(Inches(0.57), Inches(1.30), Inches(8.90), Inches(4.10))
    tf6 = tb6.text_frame
    tf6.word_wrap = True
    tf6.margin_left = tf6.margin_top = tf6.margin_right = tf6.margin_bottom = 0

    s6_content = [
        ("Week 1:", True, Pt(16), Pt(0)),
        ("●   Introduction to Software Development environment, toolchain setup, and project requirements.", False, Pt(14.5), Pt(4)),
        ("●   Setup Node.js runtime, Vite 6 bundler, and established modular project file architecture.", False, Pt(14.5), Pt(2)),
        ("●   Evaluated Indian Railways live transit APIs and analyzed rate limits and payload schemas.", False, Pt(14.5), Pt(2)),
        ("Week 2:", True, Pt(16), Pt(12)),
        ("●   Designed and developed responsive user interfaces and modern layout using CSS custom properties.", False, Pt(14.5), Pt(4)),
        ("●   Implemented the centralized Pub-Sub reactive state store (store.js) for decoupled view reactivity.", False, Pt(14.5), Pt(2)),
        ("●   Created mock JSON train schemas (12050 Gatimaan, 22436 Vande Bharat, 12952 Rajdhani) for offline testing.", False, Pt(14.5), Pt(2)),
        ("Week 3:", True, Pt(16), Pt(12)),
        ("●   Integrated MapLibre GL JS WebGL map engine with high-FPS vector basemap tiles.", False, Pt(14.5), Pt(4)),
        ("●   Developed geoService.js for Haversine distance calculations, bearing angles, and polyline interpolation.", False, Pt(14.5), Pt(2)),
        ("●   Built dynamic animated train markers with directional heading and interactive station popups.", False, Pt(14.5), Pt(2))
    ]

    for idx, (text, bold, size, before) in enumerate(s6_content):
        p = tf6.paragraphs[0] if idx == 0 else tf6.add_paragraph()
        p.text = text
        p.font.name = "Calibri"
        p.font.size = size
        p.font.bold = bold
        p.font.color.rgb = C_BLACK
        p.space_before = before

    # =========================================================================
    # SLIDE 7: Week Wise Schedule (Weeks 4 to 6)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_white_bg(s7)
    add_standard_header(s7, "Week Wise Schedule")

    tb7 = s7.shapes.add_textbox(Inches(0.57), Inches(1.30), Inches(8.90), Inches(4.10))
    tf7 = tb7.text_frame
    tf7.word_wrap = True
    tf7.margin_left = tf7.margin_top = tf7.margin_right = tf7.margin_bottom = 0

    s7_content = [
        ("Week 4:", True, Pt(16), Pt(0)),
        ("●   Configured Vite development reverse proxy to eliminate browser Cross-Origin Resource Sharing (CORS) blocks.", False, Pt(14.5), Pt(4)),
        ("●   Engineered resilient polling service (railradarService.js) with cache TTL and seamless local mock fallback.", False, Pt(14.5), Pt(2)),
        ("●   Integrated OpenWeatherMap and OpenTopography APIs for coordinate-level weather and SRTM terrain elevation.", False, Pt(14.5), Pt(2)),
        ("Week 5:", True, Pt(16), Pt(12)),
        ("●   Developed the Live Station Board module simulating digital railway junction split-flap departure/arrival displays.", False, Pt(14.5), Pt(4)),
        ("●   Built the HTML5 Canvas Real-Time Speed & Elevation Profile telemetry graph with Retina display scaling.", False, Pt(14.5), Pt(2)),
        ("●   Implemented dark mode semantic tokens, slide-in Favourites drawer, and native browser push notifications for delays.", False, Pt(14.5), Pt(2)),
        ("Week 6:", True, Pt(16), Pt(12)),
        ("●   Implemented Progressive Web Application (PWA) service worker (sw.js) and manifest for offline execution.", False, Pt(14.5), Pt(4)),
        ("●   Completed testing, debugging, and application optimization (achieved 98/100 Lighthouse performance score).", False, Pt(14.5), Pt(2)),
        ("●   Prepared project documentation, technical ITS report, and final presentation.", False, Pt(14.5), Pt(2))
    ]

    for idx, (text, bold, size, before) in enumerate(s7_content):
        p = tf7.paragraphs[0] if idx == 0 else tf7.add_paragraph()
        p.text = text
        p.font.name = "Calibri"
        p.font.size = size
        p.font.bold = bold
        p.font.color.rgb = C_BLACK
        p.space_before = before

    # =========================================================================
    # SLIDE 8: Technology Used
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_white_bg(s8)
    add_standard_header(s8, "Technology Used")

    tb8 = s8.shapes.add_textbox(Inches(0.57), Inches(1.35), Inches(8.90), Inches(4.00))
    tf8 = tb8.text_frame
    tf8.word_wrap = True
    tf8.margin_left = tf8.margin_top = tf8.margin_right = tf8.margin_bottom = 0

    s8_bullets = [
        [("●   Core Web Stack – ", True), ("HTML5 Semantic Elements, Modern CSS3 Custom Properties, and Vanilla JavaScript (ES2022 Modules).", False)],
        [("●   Vector Map & Spatial Rendering – ", True), ("MapLibre GL JS 5.1 (hardware-accelerated WebGL vector cartography with CARTO basemaps).", False)],
        [("●   Build Toolchain & Dev Server – ", True), ("Vite 6 (ESM development server, Hot Module Replacement, and Rollup production bundler).", False)],
        [("●   Telemetry & Canvas Visualization – ", True), ("HTML5 2D Canvas API with devicePixelRatio scaling for Retina display clarity.", False)],
        [("●   Transit & Environmental APIs – ", True), ("RailRadar Transit API, OpenWeatherMap live weather API, and OpenTopography SRTM DEM.", False)],
        [("●   State Architecture & PWA – ", True), ("Custom reactive Pub-Sub Store (store.js), Service Worker (sw.js), and Web App Manifest.", False)],
        [("●   Version Control & Development Tools – ", True), ("Git, GitHub, Visual Studio Code, npm, and Google Chrome DevTools.", False)]
    ]

    for idx, item in enumerate(s8_bullets):
        p = tf8.paragraphs[0] if idx == 0 else tf8.add_paragraph()
        p.space_before = Pt(10)
        for text, bold in item:
            r = p.add_run()
            r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(14.5)
            r.font.bold = bold
            r.font.color.rgb = C_BLACK

    # =========================================================================
    # SLIDE 9: RailGaadi - Project Description (Project Introduction)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_white_bg(s9)
    add_standard_header(s9, "RailGaadi - Project Description")

    tb9 = s9.shapes.add_textbox(Inches(0.57), Inches(1.30), Inches(8.90), Inches(4.10))
    tf9 = tb9.text_frame
    tf9.word_wrap = True
    tf9.margin_left = tf9.margin_top = tf9.margin_right = tf9.margin_bottom = 0

    p_sub9 = tf9.paragraphs[0]
    p_sub9.text = "Project Introduction"
    p_sub9.font.name = "Calibri"
    p_sub9.font.size = Pt(24)
    p_sub9.font.bold = True
    p_sub9.font.color.rgb = C_BLACK
    p_sub9.space_after = Pt(14)

    p9_body1 = tf9.add_paragraph()
    p9_body1.space_before = Pt(8)
    r1 = p9_body1.add_run()
    r1.text = "The "
    r2 = p9_body1.add_run()
    r2.text = "RailGaadi (Train Tracker) Web Application"
    r2.font.bold = True
    r3 = p9_body1.add_run()
    r3.text = " is a high-performance, real-time transit intelligence and spatial passenger tracking platform designed to provide passengers with comprehensive, live situational awareness across the Indian Railway network. It transforms traditional static timetables into an interactive, multi-dimensional journey monitoring experience."
    for r in [r1, r2, r3]:
        r.font.name = "Calibri"
        r.font.size = Pt(15.5)
        r.font.color.rgb = C_BLACK

    p9_body2 = tf9.add_paragraph()
    p9_body2.space_before = Pt(14)
    r4 = p9_body2.add_run()
    r4.text = "The application is developed as a "
    r5 = p9_body2.add_run()
    r5.text = "lightweight, framework-free web platform"
    r5.font.bold = True
    r6 = p9_body2.add_run()
    r6.text = " utilizing modern ES Modules, HTML5 Canvas, and "
    r7 = p9_body2.add_run()
    r7.text = "MapLibre GL JS WebGL cartography"
    r7.font.bold = True
    r8 = p9_body2.add_run()
    r8.text = ", integrated with live GPS telemetry, terrain elevation models, and micro-climate forecasts. It provides features such as "
    r9 = p9_body2.add_run()
    r9.text = "live dynamic train positioning, route gradient profiling, digital station departure boards, delay status alerts, and offline PWA capability"
    r9.font.bold = True
    r10 = p9_body2.add_run()
    r10.text = ", offering travelers a seamless, reliable, and user-friendly tracking interface."
    for r in [r4, r5, r6, r7, r8, r9, r10]:
        r.font.name = "Calibri"
        r.font.size = Pt(15.5)
        r.font.color.rgb = C_BLACK

    # =========================================================================
    # SLIDE 10: RailGaadi - Project Description (Project Scope)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_white_bg(s10)
    add_standard_header(s10, "RailGaadi - Project Description")

    tb10 = s10.shapes.add_textbox(Inches(0.57), Inches(1.30), Inches(8.90), Inches(4.10))
    tf10 = tb10.text_frame
    tf10.word_wrap = True
    tf10.margin_left = tf10.margin_top = tf10.margin_right = tf10.margin_bottom = 0

    p_sub10 = tf10.paragraphs[0]
    p_sub10.text = "Project Scope"
    p_sub10.font.name = "Calibri"
    p_sub10.font.size = Pt(24)
    p_sub10.font.bold = True
    p_sub10.font.color.rgb = C_BLACK
    p_sub10.space_after = Pt(12)

    s10_bullets = [
        [("●   Build a responsive interface using ", False), ("modern Vanilla JavaScript", True), (" with structured routing and reusable components.", False)],
        [("●   Enable users to search and track trains in real-time with ", False), ("smooth spatial polyline rendering and live speed telemetry", True), (".", False)],
        [("●   Implement an authenticated local reverse proxy to bypass ", False), ("browser CORS constraints", True), (" and ensure reliable API communication.", False)],
        [("●   Integrate multi-source external APIs including ", False), ("RailRadar transit telemetry, OpenWeatherMap conditions, and SRTM elevation", True), (".", False)],
        [("●   Engineer defensive offline resilience through ", False), ("exponential backoff retries, local cached fallbacks, and PWA service workers", True), (".", False)],
        [("●   Develop interactive station intelligence features including ", False), ("digital live arrival/departure flipboards and filterable delay analytics", True), (".", False)]
    ]

    for item in s10_bullets:
        p = tf10.add_paragraph()
        p.space_before = Pt(10)
        for text, bold in item:
            r = p.add_run()
            r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(15.0)
            r.font.bold = bold
            r.font.color.rgb = C_BLACK

    # =========================================================================
    # SLIDE 11: RailGaadi - Project Description (Module Assigned)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_white_bg(s11)
    add_standard_header(s11, "RailGaadi - Project Description")

    tb11 = s11.shapes.add_textbox(Inches(0.57), Inches(1.30), Inches(8.90), Inches(4.10))
    tf11 = tb11.text_frame
    tf11.word_wrap = True
    tf11.margin_left = tf11.margin_top = tf11.margin_right = tf11.margin_bottom = 0

    p_sub11 = tf11.paragraphs[0]
    p_sub11.text = "Module Assigned"
    p_sub11.font.name = "Calibri"
    p_sub11.font.size = Pt(24)
    p_sub11.font.bold = True
    p_sub11.font.color.rgb = C_BLACK
    p_sub11.space_after = Pt(12)

    s11_bullets = [
        [("●   Spatial Cartography & Trajectory Module – ", True), ("Implemented MapLibre GL vector map integration, train heading interpolation, and station markers.", False)],
        [("●   Reactive State Machine (store.js) – ", True), ("Architected a centralized Pub-Sub state store coordinating telemetry updates across decoupled views.", False)],
        [("●   Resilient Transit API Service (railradarService.js) – ", True), ("Engineered network fetch mechanisms with TTL caching, CORS proxying, and automated mock failover.", False)],
        [("●   Speed Telemetry & Elevation Canvas (analyticsView.js) – ", True), ("Developed double-buffered HTML5 canvas graphics with Retina devicePixelRatio scaling.", False)],
        [("●   Electronic Station Board (stationBoard.js) – ", True), ("Created live digital split-flap display boards with arrival/departure filtering and real-time delay flags.", False)],
        [("●   Offline PWA & Notification Engine (sw.js, notificationService.js) – ", True), ("Implemented service worker caching and native browser push notifications for train delays.", False)]
    ]

    for item in s11_bullets:
        p = tf11.add_paragraph()
        p.space_before = Pt(10)
        for text, bold in item:
            r = p.add_run()
            r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(14.5)
            r.font.bold = bold
            r.font.color.rgb = C_BLACK

    # =========================================================================
    # SLIDE 12: Learning
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_white_bg(s12)
    add_standard_header(s12, "Learning")

    tb12 = s12.shapes.add_textbox(Inches(0.57), Inches(1.35), Inches(8.90), Inches(4.00))
    tf12 = tb12.text_frame
    tf12.word_wrap = True
    tf12.margin_left = tf12.margin_top = tf12.margin_right = tf12.margin_bottom = 0

    s12_bullets = [
        [("●   Improved skills in ", False), ("HTML5, CSS3, modern JavaScript (ES2022+), and modular architecture", True), (".", False)],
        [("●   Gained deep practical experience in ", False), ("GIS mathematics, spherical trigonometry (Haversine formula), and polyline interpolation", True), (".", False)],
        [("●   Learned to use ", False), ("Git & GitHub", True), (" for production version control, branching workflows, and collaboration.", False)],
        [("●   Understood ", False), ("API integration and reverse proxying", True), (" for robust communication between frontend and upstream transit services.", False)],
        [("●   Gained comprehensive knowledge of ", False), ("MapLibre GL WebGL mapping, HTML5 Canvas graphics, and Retina display scaling", True), (".", False)],
        [("●   Mastered building ", False), ("Progressive Web Applications (PWA)", True), (" with service worker offline caching and native push notifications.", False)],
        [("●   Enhanced ", False), ("debugging, problem-solving, teamwork, and technical presentation skills", True), (".", False)]
    ]

    for idx, item in enumerate(s12_bullets):
        p = tf12.paragraphs[0] if idx == 0 else tf12.add_paragraph()
        p.space_before = Pt(10)
        for text, bold in item:
            r = p.add_run()
            r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(14.5)
            r.font.bold = bold
            r.font.color.rgb = C_BLACK

    # =========================================================================
    # SLIDE 13: Snapshots of Project
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_white_bg(s13)
    add_standard_header(s13, "Snapshots of Project")

    # Add 4 snapshot panels matching original layout
    if os.path.exists(snap_map):
        s13.shapes.add_picture(snap_map, Inches(0.57), Inches(1.30), Inches(4.35), Inches(1.95))
    if os.path.exists(snap_board):
        s13.shapes.add_picture(snap_board, Inches(5.08), Inches(1.30), Inches(4.35), Inches(1.95))
    if os.path.exists(snap_analytics):
        s13.shapes.add_picture(snap_analytics, Inches(0.57), Inches(3.40), Inches(4.35), Inches(1.95))
    if os.path.exists(snap_terminal):
        s13.shapes.add_picture(snap_terminal, Inches(5.08), Inches(3.40), Inches(4.35), Inches(1.95))

    # =========================================================================
    # SLIDE 14: Important Code Snippets (Requested by User!)
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_white_bg(s14)
    add_standard_header(s14, "Important Code Snippets")

    # Code Box Left: Haversine Math & Pub/Sub Store
    tb_c_left = s14.shapes.add_textbox(Inches(0.57), Inches(1.25), Inches(4.35), Inches(4.15))
    tf_c_left = tb_c_left.text_frame
    tf_c_left.word_wrap = True
    tf_c_left.margin_left = tf_c_left.margin_top = tf_c_left.margin_right = tf_c_left.margin_bottom = 0

    p_cl_head = tf_c_left.paragraphs[0]
    p_cl_head.text = "1. Spherical Haversine Distance (geoService.js)"
    p_cl_head.font.name = "Calibri"
    p_cl_head.font.size = Pt(13)
    p_cl_head.font.bold = True
    p_cl_head.font.color.rgb = C_BLACK

    p_code1 = tf_c_left.add_paragraph()
    p_code1.text = (
        "static distanceKm(lat1, lon1, lat2, lon2) {\n"
        "  const R = 6371; // Earth radius in km\n"
        "  const dLat = (lat2 - lat1) * Math.PI / 180;\n"
        "  const dLon = (lon2 - lon1) * Math.PI / 180;\n"
        "  const a = Math.sin(dLat/2)**2 +\n"
        "            Math.cos(lat1*Math.PI/180) *\n"
        "            Math.cos(lat2*Math.PI/180) * Math.sin(dLon/2)**2;\n"
        "  return R * 2 * Math.atan2(Math.sqrt(a), Math.sqrt(1 - a));\n"
        "}"
    )
    p_code1.font.name = "Courier New"
    p_code1.font.size = Pt(9.5)
    p_code1.font.color.rgb = C_BLACK
    p_code1.space_before = Pt(3)

    p_cl_head2 = tf_c_left.add_paragraph()
    p_cl_head2.text = "2. Pub/Sub Reactive State Store (store.js)"
    p_cl_head2.font.name = "Calibri"
    p_cl_head2.font.size = Pt(13)
    p_cl_head2.font.bold = True
    p_cl_head2.font.color.rgb = C_BLACK
    p_cl_head2.space_before = Pt(12)

    p_code2 = tf_c_left.add_paragraph()
    p_code2.text = (
        "subscribe(listener) {\n"
        "  this.listeners.add(listener);\n"
        "  return () => this.listeners.delete(listener);\n"
        "}\n"
        "setState(updates) {\n"
        "  this.state = { ...this.state, ...updates };\n"
        "  const changed = Object.keys(updates);\n"
        "  this.listeners.forEach(fn => fn(this.state, changed));\n"
        "}"
    )
    p_code2.font.name = "Courier New"
    p_code2.font.size = Pt(9.5)
    p_code2.font.color.rgb = C_BLACK
    p_code2.space_before = Pt(3)

    # Code Box Right: Retina Canvas & Reverse Proxy
    tb_c_right = s14.shapes.add_textbox(Inches(5.08), Inches(1.25), Inches(4.35), Inches(4.15))
    tf_c_right = tb_c_right.text_frame
    tf_c_right.word_wrap = True
    tf_c_right.margin_left = tf_c_right.margin_top = tf_c_right.margin_right = tf_c_right.margin_bottom = 0

    p_cr_head = tf_c_right.paragraphs[0]
    p_cr_head.text = "3. Retina Display Canvas Scaling (analyticsView.js)"
    p_cr_head.font.name = "Calibri"
    p_cr_head.font.size = Pt(13)
    p_cr_head.font.bold = True
    p_cr_head.font.color.rgb = C_BLACK

    p_code3 = tf_c_right.add_paragraph()
    p_code3.text = (
        "const ctx = canvas.getContext('2d');\n"
        "const dpr = window.devicePixelRatio || 1;\n"
        "const rect = canvas.getBoundingClientRect();\n"
        "canvas.width = rect.width * dpr;\n"
        "canvas.height = rect.height * dpr;\n"
        "ctx.scale(dpr, dpr); // Crystal-clear on 4K / Retina\n"
        "ctx.clearRect(0, 0, rect.width, rect.height);"
    )
    p_code3.font.name = "Courier New"
    p_code3.font.size = Pt(9.5)
    p_code3.font.color.rgb = C_BLACK
    p_code3.space_before = Pt(3)

    p_cr_head2 = tf_c_right.add_paragraph()
    p_cr_head2.text = "4. Reverse Proxy CORS Bypass (vite.config.ts)"
    p_cr_head2.font.name = "Calibri"
    p_cr_head2.font.size = Pt(13)
    p_cr_head2.font.bold = True
    p_cr_head2.font.color.rgb = C_BLACK
    p_cr_head2.space_before = Pt(12)

    p_code4 = tf_c_right.add_paragraph()
    p_code4.text = (
        "defineConfig({\n"
        "  server: {\n"
        "    port: 8080,\n"
        "    proxy: {\n"
        "      '/api/railradar': {\n"
        "        target: 'https://api.railradar.in/v1',\n"
        "        changeOrigin: true,\n"
        "        rewrite: p => p.replace(/^\\/api\\/railradar/, '')\n"
        "      }\n"
        "    }\n"
        "  }\n"
        "});"
    )
    p_code4.font.name = "Courier New"
    p_code4.font.size = Pt(9.5)
    p_code4.font.color.rgb = C_BLACK
    p_code4.space_before = Pt(3)

    # =========================================================================
    # SLIDE 15: Future Scope
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_white_bg(s15)
    add_standard_header(s15, "Future Scope")

    tb15 = s15.shapes.add_textbox(Inches(0.57), Inches(1.35), Inches(8.90), Inches(4.00))
    tf15 = tb15.text_frame
    tf15.word_wrap = True
    tf15.margin_left = tf15.margin_top = tf15.margin_right = tf15.margin_bottom = 0

    s15_bullets = [
        [("●   Crowdsourced Live Tracking – ", True), ("Allow on-board passengers to validate real-time GPS locations and coach occupancy.", False)],
        [("●   AI Predictive Delay Modeling – ", True), ("Integrate machine learning to predict junction bottlenecks and delay propagation across sectors.", False)],
        [("●   Multi-Modal Transit Integration – ", True), ("Connect Indian Railways schedules with metro lines and local buses for unified journey planning.", False)],
        [("●   WebSocket & MQTT Streaming – ", True), ("Upgrade polling intervals to low-latency bi-directional streams for sub-second updates.", False)],
        [("●   Cross-Platform Mobile App – ", True), ("Extend the web platform into native Android and iOS applications with offline background notifications.", False)],
        [("●   Offline Vector Map Caching – ", True), ("Pre-cache regional map boundaries locally so tracks and terrain render inside isolated tunnels.", False)]
    ]

    for idx, item in enumerate(s15_bullets):
        p = tf15.paragraphs[0] if idx == 0 else tf15.add_paragraph()
        p.space_before = Pt(10)
        for text, bold in item:
            r = p.add_run()
            r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(14.5)
            r.font.bold = bold
            r.font.color.rgb = C_BLACK

    # =========================================================================
    # SLIDE 16: Conclusion
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    set_white_bg(s16)
    add_standard_header(s16, "Conclusion")

    tb16 = s16.shapes.add_textbox(Inches(0.57), Inches(1.35), Inches(8.90), Inches(4.00))
    tf16 = tb16.text_frame
    tf16.word_wrap = True
    tf16.margin_left = tf16.margin_top = tf16.margin_right = tf16.margin_bottom = 0

    s16_bullets = [
        [("●   Gained practical experience in ", False), ("Software Development", True), (" through real-world project implementation.", False)],
        [("●   Strengthened knowledge of ", False), ("HTML, CSS, modern JavaScript (ES Modules), MapLibre GL, and API integration", True), (".", False)],
        [("●   Successfully developed the ", False), ("RailGaadi (Train Tracker) Web Application", True), (" with live telemetry, elevation curves, and station boards.", False)],
        [("●   Learned to use professional tools and technologies like ", False), ("Git, GitHub, VS Code, Vite, and PWA Service Workers", True), (" for development.", False)],
        [("●   Enhanced ", False), ("teamwork, communication, problem-solving, and time management", True), (" in a professional environment.", False)],
        [("●   Overall, the internship boosted my confidence and prepared me for ", False), ("future opportunities in full-stack web and software engineering", True), (".", False)]
    ]

    for idx, item in enumerate(s16_bullets):
        p = tf16.paragraphs[0] if idx == 0 else tf16.add_paragraph()
        p.space_before = Pt(10)
        for text, bold in item:
            r = p.add_run()
            r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(14.5)
            r.font.bold = bold
            r.font.color.rgb = C_BLACK

    # =========================================================================
    # SLIDE 17: References
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    set_white_bg(s17)
    add_standard_header(s17, "References")

    tb17 = s17.shapes.add_textbox(Inches(0.57), Inches(1.30), Inches(8.90), Inches(4.10))
    tf17 = tb17.text_frame
    tf17.word_wrap = True
    tf17.margin_left = tf17.margin_top = tf17.margin_right = tf17.margin_bottom = 0

    s17_items = [
        [("●   Sumit DigiTech Private Limited – ", True), ("Official Website and Training Resources", False)],
        [("●   MapLibre GL JS Documentation – ", True), ("https://maplibre.org/maplibre-gl-js/docs/", False)],
        [("●   MDN Web Docs – ", True), ("HTML, CSS, JavaScript, Canvas API & Service Workers", False)],
        [("●   Vite Official Documentation – ", True), ("Next Generation Frontend Tooling", False)],
        [("●   OpenWeatherMap & OpenTopography – ", True), ("Weather & SRTM Elevation APIs", False)],
        [("●   Git & GitHub Guides – ", True), ("Version Control & Repository Management", False)],
        [("Other References", True), ("", False)],
        [("●   Tutorials & articles from platforms like ", False), ("GeeksforGeeks, W3Schools, FreeCodeCamp", True)],
        [("●   Guidance and mentorship provided by the ", False), ("Sumit DigiTech team & Poornima Faculty Coordinators", True)]
    ]

    for idx, item in enumerate(s17_items):
        p = tf17.paragraphs[0] if idx == 0 else tf17.add_paragraph()
        if item[0][0] == "Other References":
            p.space_before = Pt(14)
            r = p.add_run()
            r.text = "Other References"
            r.font.name = "Calibri"
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = C_BLACK
        else:
            p.space_before = Pt(8)
            for text, bold in item:
                if text:
                    r = p.add_run()
                    r.text = text
                    r.font.name = "Calibri"
                    r.font.size = Pt(14.5)
                    r.font.bold = bold
                    r.font.color.rgb = C_BLACK

    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation at: {output_path}")

if __name__ == "__main__":
    out_file = "RailGaadi_ITS_Presentation_Nikhil_Kumar.pptx"
    build_presentation(out_file)
    # Also save to public for immediate download
    build_presentation(os.path.join("public", out_file))
